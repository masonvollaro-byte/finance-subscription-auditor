import io
import os
import re
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import quote

import pandas as pd
import pdfplumber
import plotly.express as px
import streamlit as st
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TRANSACTION_COLUMNS = [
    "date",
    "description",
    "amount",
    "type",
    "source",
]

CADENCES = {
    "Weekly": (6, 8, 52),
    "Monthly": (27, 32, 12),
    "Quarterly": (85, 95, 4),
    "Annual": (350, 380, 1),
}

KNOWN_MERCHANTS = {
    "NETFLIX": "Netflix",
    "SPOTIFY": "Spotify",
    "NEW YORK TIMES": "New York Times",
    "NYT": "New York Times",
    "HULU": "Hulu",
    "DISNEY": "Disney+",
    "AMAZON PRIME": "Amazon Prime",
    "APPLE.COM/BILL": "Apple",
    "ADOBE": "Adobe",
    "MICROSOFT": "Microsoft",
    "DROPBOX": "Dropbox",
    "YOUTUBE": "YouTube",
}

CATEGORY_RULES = {
    "Housing": [
        "RENT",
        "MORTGAGE",
        "APARTMENT",
    ],
    "Subscriptions": [
        "NETFLIX",
        "SPOTIFY",
        "HULU",
        "DISNEY",
        "ADOBE",
        "AMAZON PRIME",
        "DROPBOX",
        "YOUTUBE",
    ],
    "Utilities": [
        "ELECTRIC",
        "UTILITY",
        "WATER",
        "INTERNET",
        "PHONE",
    ],
    "Groceries": [
        "MARKET",
        "GROCERY",
        "SUPERMARKET",
        "WALMART",
    ],
    "Dining": [
        "RESTAURANT",
        "COFFEE",
        "CAFE",
        "DOORDASH",
        "UBER EATS",
    ],
    "Transportation": [
        "GAS",
        "SHELL",
        "EXXON",
        "UBER",
        "LYFT",
        "PARKING",
    ],
}

WEB_COLORS = {
    "Price decreased": "#C6EFCE",
    "Slight increase": "#FFF2CC",
    "Large increase": "#FFC7CE",
}

EXCEL_COLORS = {
    "Price decreased": "C6EFCE",
    "Slight increase": "FFF2CC",
    "Large increase": "FFC7CE",
}

DATE_PATTERN = (
    r"(?P<date>\d{1,2}[/-]\d{1,2}"
    r"(?:[/-]\d{2,4})?)"
)

AMOUNT_PATTERN = (
    r"(?P<amount>\(?-?\$?\d{1,3}"
    r"(?:,\d{3})*\.\d{2}\)?)"
)

LINE_PATTERN = re.compile(
    rf"^\s*{DATE_PATTERN}\s+"
    rf"(?P<description>.+?)\s+"
    rf"{AMOUNT_PATTERN}\s*$"
)


def parse_amount(raw_amount):
    value = (
        raw_amount
        .replace("$", "")
        .replace(",", "")
        .strip()
    )

    parentheses_negative = (
        value.startswith("(")
        and value.endswith(")")
    )

    value = value.strip("()")
    amount = float(value)

    if parentheses_negative:
        amount = -amount

    return amount


def parse_pdf(file_bytes, filename):
    rows = []

    with pdfplumber.open(
        io.BytesIO(file_bytes)
    ) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""

            for line in text.splitlines():
                match = LINE_PATTERN.match(line)

                if not match:
                    continue

                try:
                    amount = parse_amount(
                        match.group("amount")
                    )
                except ValueError:
                    continue

                rows.append(
                    {
                        "date": match.group("date"),
                        "description": match.group(
                            "description"
                        ).strip(),
                        "amount": amount,
                        "type": (
                            "credit"
                            if amount < 0
                            else "debit"
                        ),
                        "source": Path(filename).name,
                    }
                )

    return pd.DataFrame(
        rows,
        columns=TRANSACTION_COLUMNS,
    )


def clean_description(description):
    text = str(
        description or ""
    ).upper().strip()

    text = re.sub(
        r"\b(?:ACH|POS|DEBIT|PURCHASE|"
        r"CARD|RECURRING)\b",
        " ",
        text,
    )

    text = re.sub(
        r"\b\d{3}[-.) ]?\d{3}"
        r"[-. ]?\d{4}\b",
        " ",
        text,
    )

    text = re.sub(
        r"\b(?:REF|TRACE|ID)\s*#?\s*"
        r"[A-Z0-9-]+\b",
        " ",
        text,
    )

    return re.sub(
        r"\s+",
        " ",
        text,
    ).strip(" -*").title()


def normalize_merchant(description):
    cleaned = clean_description(description)
    uppercase = cleaned.upper()

    for fragment, merchant in (
        KNOWN_MERCHANTS.items()
    ):
        if fragment in uppercase:
            return merchant

    cleaned = re.sub(
        r"\s+#?\d{4,}\b.*$",
        "",
        cleaned,
    )

    return cleaned[:80] or "Unknown merchant"


def categorize_transaction(merchant):
    uppercase = str(merchant).upper()

    for category, keywords in (
        CATEGORY_RULES.items()
    ):
        if any(
            keyword in uppercase
            for keyword in keywords
        ):
            return category

    return "Other"


def clean_transactions(frame):
    columns = (
        TRANSACTION_COLUMNS
        + ["merchant", "category"]
    )

    if frame.empty:
        return pd.DataFrame(columns=columns)

    data = frame.copy()

    for column in TRANSACTION_COLUMNS:
        if column not in data:
            data[column] = ""

    data["date"] = pd.to_datetime(
        data["date"],
        errors="coerce",
    )

    data["amount"] = pd.to_numeric(
        data["amount"],
        errors="coerce",
    )

    data["description"] = (
        data["description"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    data["type"] = (
        data["type"]
        .fillna("debit")
        .astype(str)
        .str.lower()
    )

    data["source"] = (
        data["source"]
        .fillna("")
        .astype(str)
    )

    data = data.dropna(
        subset=["date", "amount"]
    )

    data = data[
        data["description"] != ""
    ]

    data["amount"] = (
        data["amount"]
        .abs()
        .round(2)
    )

    data["merchant"] = (
        data["description"].map(
            normalize_merchant
        )
    )

    data["category"] = (
        data["merchant"].map(
            categorize_transaction
        )
    )

    return (
        data[columns]
        .sort_values(["merchant", "date"])
        .reset_index(drop=True)
    )


def classify_cadence(intervals):
    for name, values in CADENCES.items():
        minimum, maximum, multiplier = values

        confidence = float(
            intervals.between(
                minimum,
                maximum,
            ).mean()
        )

        if confidence >= 0.60:
            return name, confidence, multiplier

    return None, 0.0, 0


def classify_price_status(percent):
    if percent < 0:
        return "Price decreased"

    if 1 <= percent < 20:
        return "Slight increase"

    if percent >= 20:
        return "Large increase"

    return "No significant change"


def detect_recurring(transactions):
    findings = []

    debits = transactions[
        transactions["type"] != "credit"
    ].copy()

    for merchant, group in (
        debits.groupby("merchant")
    ):
        group = (
            group
            .sort_values("date")
            .drop_duplicates(
                subset=["date", "amount"]
            )
        )

        if len(group) < 2:
            continue

        intervals = (
            group["date"]
            .diff()
            .dt.days
            .dropna()
        )

        cadence, confidence, multiplier = (
            classify_cadence(intervals)
        )

        if cadence is None:
            continue

        latest = float(
            group["amount"].iloc[-1]
        )

        previous = float(
            group["amount"].iloc[-2]
        )

        change = round(
            latest - previous,
            2,
        )

        percent = (
            round(
                change / previous * 100,
                1,
            )
            if previous
            else 0.0
        )

        latest_date = (
            group["date"].iloc[-1]
        )

        median_interval = max(
            1,
            int(round(intervals.median())),
        )

        findings.append(
            {
                "merchant": merchant,
                "cadence": cadence,
                "confidence": round(
                    confidence * 100
                ),
                "occurrences": len(group),
                "latest_amount": latest,
                "previous_amount": previous,
                "price_change": change,
                "price_change_percent": percent,
                "price_status":
                    classify_price_status(
                        percent
                    ),
                "latest_charge":
                    latest_date,
                "estimated_next_charge":
                    latest_date
                    + pd.Timedelta(
                        days=median_interval
                    ),
                "estimated_annual_cost":
                    round(
                        latest * multiplier,
                        2,
                    ),
            }
        )

    result = pd.DataFrame(findings)

    if result.empty:
        return result

    return result.sort_values(
        [
            "price_change_percent",
            "estimated_annual_cost",
        ],
        ascending=[False, False],
    ).reset_index(drop=True)


def add_transaction_price_status(
    transactions,
    recurring,
):
    result = transactions.copy()

    result["price_status"] = (
        "No significant change"
    )

    transaction_dates = pd.to_datetime(
        result["date"]
    ).dt.normalize()

    for row in recurring.itertuples(
        index=False
    ):
        latest_date = pd.Timestamp(
            row.latest_charge
        ).normalize()

        matching = (
            (result["merchant"] == row.merchant)
            & (transaction_dates == latest_date)
        )

        result.loc[
            matching,
            "price_status",
        ] = row.price_status

    return result


def style_price_rows(dataframe):
    def style_row(row):
        status = row.get(
            "price_status",
            "No significant change",
        )

        color = WEB_COLORS.get(status)

        if not color:
            return ["" for _ in row]

        return [
            (
                f"background-color: {color}; "
                "color: black"
            )
            for _ in row
        ]

    return dataframe.style.apply(
        style_row,
        axis=1,
    )


def escape_calendar_text(value):
    return (
        str(value)
        .replace("\\", "\\\\")
        .replace(",", "\\,")
        .replace(";", "\\;")
        .replace("\n", "\\n")
    )


def create_calendar_file(recurring):
    timestamp = datetime.now(
        timezone.utc
    ).strftime("%Y%m%dT%H%M%SZ")

    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        (
            "PRODID:-//Finance Subscription "
            "Auditor//EN"
        ),
    ]

    for row in recurring.itertuples(
        index=False
    ):
        start = pd.Timestamp(
            row.estimated_next_charge
        ).to_pydatetime()

        end = start + timedelta(days=1)

        merchant = escape_calendar_text(
            row.merchant
        )

        lines.extend(
            [
                "BEGIN:VEVENT",
                f"UID:{uuid.uuid4()}",
                f"DTSTAMP:{timestamp}",
                (
                    "DTSTART;VALUE=DATE:"
                    f"{start:%Y%m%d}"
                ),
                (
                    "DTEND;VALUE=DATE:"
                    f"{end:%Y%m%d}"
                ),
                (
                    f"SUMMARY:{merchant} "
                    f"charge ${row.latest_amount:.2f}"
                ),
                (
                    "DESCRIPTION:Review or cancel "
                    "this recurring charge."
                ),
                "BEGIN:VALARM",
                "TRIGGER:-P3D",
                "ACTION:DISPLAY",
                (
                    f"DESCRIPTION:Review "
                    f"{merchant} subscription"
                ),
                "END:VALARM",
                "END:VEVENT",
            ]
        )

    lines.extend(
        ["END:VCALENDAR", ""]
    )

    return "\r\n".join(
        lines
    ).encode("utf-8")


def build_email_summary(
    recurring,
    transactions,
):
    annual_total = recurring[
        "estimated_annual_cost"
    ].sum()

    lines = [
        "Finance & Subscription Auditor",
        "",
        (
            "Transactions analyzed: "
            f"{len(transactions)}"
        ),
        (
            "Recurring charges found: "
            f"{len(recurring)}"
        ),
        (
            "Estimated monthly cost: "
            f"${annual_total / 12:,.2f}"
        ),
        (
            "Estimated annual cost: "
            f"${annual_total:,.2f}"
        ),
        "",
        "Recurring charges:",
    ]

    for row in recurring.itertuples(
        index=False
    ):
        lines.append(
            f"- {row.merchant}: "
            f"${row.latest_amount:.2f}, "
            f"{row.cadence}, "
            f"{row.price_status}"
        )

    lines.extend(
        [
            "",
            (
                "Verify this report against "
                "the original statements."
            ),
        ]
    )

    return "\n".join(lines)


def create_email_link(
    recipient,
    subject,
    body,
):
    return (
        f"mailto:{quote(recipient.strip())}"
        f"?subject={quote(subject)}"
        f"&body={quote(body)}"
    )


def generate_ai_draft(
    merchant,
    latest_amount,
    previous_amount,
    cadence,
    action,
):
    try:
        from openai import OpenAI
    except ImportError as error:
        raise RuntimeError(
            "The OpenAI package is not installed."
        ) from error

    if not os.getenv("OPENAI_API_KEY"):
        raise RuntimeError(
            "OPENAI_API_KEY is not configured."
        )

    client = OpenAI()

    context = f"""
Merchant: {merchant}
Latest charge: ${latest_amount:.2f}
Previous charge: ${previous_amount:.2f}
Billing cadence: {cadence}
Requested document: {action}
"""

    response = client.responses.create(
        model=os.getenv(
            "OPENAI_MODEL",
            "gpt-5.4-nano",
        ),
        instructions=(
            "Write a concise consumer billing "
            "message. Do not invent account "
            "numbers, dates, laws, refund rights, "
            "or contact information. Use "
            "placeholders where needed."
        ),
        input=context,
        store=False,
    )

    return response.output_text


def make_excel_value(value):
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()

    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            return value

    return value


def format_excel_sheet(
    sheet,
    dataframe,
):
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="17365D",
    )

    stripe_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    sheet.sheet_view.showGridLines = False

    for column_number, column_name in enumerate(
        dataframe.columns,
        start=1,
    ):
        cell = sheet.cell(
            row=1,
            column=column_number,
            value=str(column_name)
            .replace("_", " ")
            .title(),
        )

        cell.fill = header_fill

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row_number, row in enumerate(
        dataframe.itertuples(
            index=False,
            name=None,
        ),
        start=2,
    ):
        for column_number, value in enumerate(
            row,
            start=1,
        ):
            cell = sheet.cell(
                row=row_number,
                column=column_number,
                value=make_excel_value(value),
            )

            cell.alignment = Alignment(
                vertical="center",
            )

            if row_number % 2 == 0:
                cell.fill = stripe_fill

    if "price_status" in dataframe.columns:
        status_column = (
            list(dataframe.columns).index(
                "price_status"
            )
            + 1
        )

        for row_number in range(
            2,
            len(dataframe) + 2,
        ):
            status = sheet.cell(
                row=row_number,
                column=status_column,
            ).value

            color = EXCEL_COLORS.get(status)

            if color:
                status_fill = PatternFill(
                    fill_type="solid",
                    fgColor=color,
                )

                for column_number in range(
                    1,
                    len(dataframe.columns) + 1,
                ):
                    sheet.cell(
                        row=row_number,
                        column=column_number,
                    ).fill = status_fill

    for column_number, column_name in enumerate(
        dataframe.columns,
        start=1,
    ):
        maximum_length = len(
            str(column_name)
        )

        for row_number in range(
            2,
            len(dataframe) + 2,
        ):
            value = sheet.cell(
                row=row_number,
                column=column_number,
            ).value

            maximum_length = max(
                maximum_length,
                len(str(value or "")),
            )

        sheet.column_dimensions[
            get_column_letter(column_number)
        ].width = min(
            maximum_length + 3,
            45,
        )

        if (
            "amount" in column_name
            or "cost" in column_name
            or column_name == "price_change"
        ):
            for row_number in range(
                2,
                len(dataframe) + 2,
            ):
                sheet.cell(
                    row=row_number,
                    column=column_number,
                ).number_format = "$#,##0.00"

        if (
            "date" in column_name
            or "charge" in column_name
        ):
            for row_number in range(
                2,
                len(dataframe) + 2,
            ):
                sheet.cell(
                    row=row_number,
                    column=column_number,
                ).number_format = "mm/dd/yyyy"

        if column_name == "price_change_percent":
            for row_number in range(
                2,
                len(dataframe) + 2,
            ):
                sheet.cell(
                    row=row_number,
                    column=column_number,
                ).number_format = '0.0"%"'

    if len(dataframe.columns) > 0:
        final_column = get_column_letter(
            len(dataframe.columns)
        )

        final_row = max(
            len(dataframe) + 1,
            1,
        )

        sheet.auto_filter.ref = (
            f"A1:{final_column}{final_row}"
        )

    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 24


def create_excel_report(
    transactions,
    recurring,
):
    workbook = Workbook()
    dashboard = workbook.active
    dashboard.title = "Dashboard"
    dashboard.sheet_view.showGridLines = False

    dashboard.merge_cells("A1:D1")
    dashboard["A1"] = (
        "Finance & Subscription Auditor"
    )

    dashboard["A1"].font = Font(
        bold=True,
        size=18,
        color="FFFFFF",
    )

    dashboard["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="17365D",
    )

    dashboard["A1"].alignment = Alignment(
        horizontal="center",
    )

    annual_total = recurring[
        "estimated_annual_cost"
    ].sum()

    metrics = [
        (
            "Transactions analyzed",
            len(transactions),
        ),
        (
            "Recurring charges",
            len(recurring),
        ),
        (
            "Monthly recurring estimate",
            annual_total / 12,
        ),
        (
            "Annual recurring estimate",
            annual_total,
        ),
        (
            "Price decreases",
            int(
                (
                    recurring["price_status"]
                    == "Price decreased"
                ).sum()
            ),
        ),
        (
            "Slight increases",
            int(
                (
                    recurring["price_status"]
                    == "Slight increase"
                ).sum()
            ),
        ),
        (
            "Large increases",
            int(
                (
                    recurring["price_status"]
                    == "Large increase"
                ).sum()
            ),
        ),
    ]

    dashboard["A3"] = "Metric"
    dashboard["B3"] = "Value"

    for cell_name in ("A3", "B3"):
        dashboard[cell_name].font = Font(
            bold=True,
            color="FFFFFF",
        )

        dashboard[cell_name].fill = PatternFill(
            fill_type="solid",
            fgColor="4472C4",
        )

    for row_number, metric in enumerate(
        metrics,
        start=4,
    ):
        dashboard.cell(
            row=row_number,
            column=1,
            value=metric[0],
        )

        dashboard.cell(
            row=row_number,
            column=2,
            value=metric[1],
        )

    dashboard["B6"].number_format = "$#,##0.00"
    dashboard["B7"].number_format = "$#,##0.00"

    category_totals = (
        transactions[
            transactions["type"] != "credit"
        ]
        .groupby(
            "category",
            as_index=False,
        )["amount"]
        .sum()
        .sort_values(
            "amount",
            ascending=False,
        )
    )

    category_start = 13

    dashboard.cell(
        row=category_start,
        column=1,
        value="Category",
    )

    dashboard.cell(
        row=category_start,
        column=2,
        value="Amount",
    )

    for cell in (
        dashboard.cell(category_start, 1),
        dashboard.cell(category_start, 2),
    ):
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="70AD47",
        )

    for row_number, row in enumerate(
        category_totals.itertuples(
            index=False
        ),
        start=category_start + 1,
    ):
        dashboard.cell(
            row=row_number,
            column=1,
            value=row.category,
        )

        dashboard.cell(
            row=row_number,
            column=2,
            value=float(row.amount),
        ).number_format = "$#,##0.00"

    if len(category_totals) > 0:
        final_row = (
            category_start
            + len(category_totals)
        )

        chart_data = Reference(
            dashboard,
            min_col=2,
            min_row=category_start,
            max_row=final_row,
        )

        chart_categories = Reference(
            dashboard,
            min_col=1,
            min_row=category_start + 1,
            max_row=final_row,
        )

        bar_chart = BarChart()
        bar_chart.title = "Spending Allocation"
        bar_chart.style = 10
        bar_chart.height = 8
        bar_chart.width = 14

        bar_chart.add_data(
            chart_data,
            titles_from_data=True,
        )

        bar_chart.set_categories(
            chart_categories
        )

        dashboard.add_chart(
            bar_chart,
            "D3",
        )

        pie_chart = PieChart()
        pie_chart.title = (
            "Spending by Category"
        )
        pie_chart.height = 8
        pie_chart.width = 12

        pie_chart.add_data(
            chart_data,
            titles_from_data=True,
        )

        pie_chart.set_categories(
            chart_categories
        )

        dashboard.add_chart(
            pie_chart,
            "D19",
        )

    dashboard.column_dimensions["A"].width = 32
    dashboard.column_dimensions["B"].width = 20
    dashboard.freeze_panes = "A3"

    transaction_sheet = (
        workbook.create_sheet(
            "Transactions"
        )
    )

    format_excel_sheet(
        transaction_sheet,
        transactions,
    )

    subscription_sheet = (
        workbook.create_sheet(
            "Subscriptions"
        )
    )

    format_excel_sheet(
        subscription_sheet,
        recurring,
    )

    confidence_sheet = workbook.create_sheet(
        "Confidence Guide"
    )

    confidence_sheet.sheet_view.showGridLines = False
    confidence_sheet.column_dimensions["A"].width = 20
    confidence_sheet.column_dimensions["B"].width = 75

    confidence_sheet["A1"] = "Confidence Level Guide"
    confidence_sheet["A1"].font = Font(
        bold=True,
        size=16,
        color="FFFFFF",
    )
    confidence_sheet["A1"].fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    confidence_sheet.merge_cells("A1:B1")

    guide_rows = [
        ("Confidence", "Meaning"),
        (
            "100%",
            "Every transaction interval matches the detected schedule.",
        ),
        (
            "80%",
            "Most transaction intervals match the detected schedule.",
        ),
        (
            "60%",
            "The minimum confidence accepted by the app.",
        ),
        (
            "Below 60%",
            "The charge is not classified as recurring.",
        ),
    ]

    for row_number, row_values in enumerate(
        guide_rows,
        start=3,
    ):
        for column_number, value in enumerate(
            row_values,
            start=1,
        ):
            cell = confidence_sheet.cell(
                row=row_number,
                column=column_number,
                value=value,
            )
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            if row_number == 3:
                cell.font = Font(
                    bold=True,
                    color="FFFFFF",
                )
                cell.fill = PatternFill(
                    "solid",
                    fgColor="4472C4",
                )

    confidence_sheet["A9"] = "Important"
    confidence_sheet["A9"].font = Font(bold=True)
    confidence_sheet["B9"] = (
        "Confidence measures how consistently a charge follows a "
        "schedule. It does not guarantee that the charge is a "
        "subscription."
    )
    confidence_sheet["B9"].alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )
    confidence_sheet.freeze_panes = "A3"

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    return output.getvalue()


def run_app():
    st.set_page_config(
        page_title=(
            "Finance & Subscription Auditor"
        ),
        page_icon="🔎",
        layout="wide",
    )

    st.title(
        "Finance & Subscription Auditor"
    )

    st.caption(
        "Privacy-first local analysis "
        "of PDF bank statements"
    )

    with st.sidebar:
        st.header("Privacy")

        st.success(
            "PDF extraction and calculations "
            "run locally."
        )

        enable_ai = st.checkbox(
            "Enable optional AI drafts",
            value=False,
        )

    uploaded_files = st.file_uploader(
        "Upload searchable PDF statements",
        type=["pdf"],
        accept_multiple_files=True,
        help=(
            "Upload at least two months for "
            "recurring-charge detection."
        ),
    )

    if not uploaded_files:
        st.info(
            "Upload two synthetic or redacted "
            "statements to begin."
        )
        return

    frames = []

    with st.status(
        "Extracting transactions...",
        expanded=True,
    ) as status:
        for uploaded_file in uploaded_files:
            try:
                frame = parse_pdf(
                    uploaded_file.getvalue(),
                    uploaded_file.name,
                )

                frames.append(frame)

                st.write(
                    f"{uploaded_file.name}: "
                    f"{len(frame)} rows found"
                )

            except Exception as error:
                st.warning(
                    f"{uploaded_file.name}: "
                    f"{error}"
                )

        status.update(
            label="Extraction complete",
            state="complete",
        )

    if not frames:
        st.error(
            "No statements could be processed."
        )
        return

    transactions = clean_transactions(
        pd.concat(
            frames,
            ignore_index=True,
        )
    )

    if transactions.empty:
        st.error(
            "No transactions were recognized. "
            "The PDFs may be scanned or use an "
            "unsupported layout."
        )
        return

    st.subheader(
        "1. Verify extracted transactions"
    )

    edited_transactions = st.data_editor(
        transactions,
        hide_index=True,
        use_container_width=True,
        num_rows="dynamic",
        column_config={
            "date":
                st.column_config.DateColumn(
                    "Date"
                ),
            "amount":
                st.column_config.NumberColumn(
                    "Amount",
                    format="$%.2f",
                ),
        },
    )

    edited_transactions = clean_transactions(
        edited_transactions
    )

    recurring = detect_recurring(
        edited_transactions
    )

    st.subheader(
        "2. Subscription audit"
    )

    if recurring.empty:
        st.warning(
            "No recurring charges were found. "
            "Try uploading additional months."
        )
        return

    annual_total = recurring[
        "estimated_annual_cost"
    ].sum()

    metric_columns = st.columns(4)

    metric_columns[0].metric(
        "Recurring charges",
        len(recurring),
    )

    metric_columns[1].metric(
        "Monthly equivalent",
        f"${annual_total / 12:,.2f}",
    )

    metric_columns[2].metric(
        "Annual cost",
        f"${annual_total:,.2f}",
    )

    metric_columns[3].metric(
        "Large increases",
        int(
            (
                recurring["price_status"]
                == "Large increase"
            ).sum()
        ),
    )

    st.caption(
        "Green = decreased, yellow = "
        "1%–19% increase, red = "
        "20% or greater increase."
    )

    with st.expander(
        "What does the confidence level mean?"
    ):
        st.markdown(
            """
- **100%:** Every transaction interval matches the detected schedule.
- **80%:** Most intervals match the detected schedule.
- **60%:** This is the minimum confidence accepted by the app.
- **Below 60%:** The charge is not classified as recurring.

Confidence measures how consistently the charge follows a schedule. It does not guarantee that the charge is a subscription.
"""
        )

    st.dataframe(
        style_price_rows(recurring),
        hide_index=True,
        use_container_width=True,
    )

    colored_transactions = (
        add_transaction_price_status(
            edited_transactions,
            recurring,
        )
    )

    st.subheader(
        "3. Color-coded transactions"
    )

    st.dataframe(
        style_price_rows(
            colored_transactions
        ),
        hide_index=True,
        use_container_width=True,
    )

    debits = colored_transactions[
        colored_transactions["type"]
        != "credit"
    ].copy()

    debits["month"] = (
        debits["date"]
        .dt.to_period("M")
        .astype(str)
    )

    latest_month = debits[
        "month"
    ].max()

    monthly_categories = (
        debits[
            debits["month"]
            == latest_month
        ]
        .groupby(
            "category",
            as_index=False,
        )["amount"]
        .sum()
    )

    st.subheader(
        "4. Monthly spending"
    )

    pie_chart = px.pie(
        monthly_categories,
        names="category",
        values="amount",
        title=(
            f"Monthly spending by category "
            f"- {latest_month}"
        ),
        hole=0.35,
    )

    st.plotly_chart(
        pie_chart,
        use_container_width=True,
    )

    allocation = (
        debits
        .groupby(
            "category",
            as_index=False,
        )["amount"]
        .sum()
        .sort_values(
            "amount",
            ascending=False,
        )
    )

    st.subheader(
        "5. Money allocation"
    )

    bar_chart = px.bar(
        allocation,
        x="category",
        y="amount",
        color="category",
        title="Where your money is allocated",
        labels={
            "category": "Category",
            "amount": "Amount spent",
        },
    )

    bar_chart.update_layout(
        showlegend=False
    )

    st.plotly_chart(
        bar_chart,
        use_container_width=True,
    )

    st.subheader(
        "6. Download formatted report"
    )

    excel_report = create_excel_report(
        colored_transactions,
        recurring,
    )

    st.download_button(
        "Download formatted Excel report",
        data=excel_report,
        file_name=(
            "finance_auditor_report.xlsx"
        ),
        mime=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )

    st.subheader(
        "7. Create an email update"
    )

    st.info(
        "The app opens your normal email "
        "program with a completed report. "
        "Review the message and click Send."
    )

    recipient = st.text_input(
        "Recipient email address",
        placeholder="name@example.com",
    )

    email_summary = build_email_summary(
        recurring,
        colored_transactions,
    )

    if (
        recipient
        and "@" in recipient
    ):
        email_link = create_email_link(
            recipient,
            "Finance Auditor Update",
            email_summary,
        )

        st.link_button(
            "Open pre-filled email",
            email_link,
            use_container_width=True,
        )

        with st.expander(
            "Preview or copy the email"
        ):
            st.text_area(
                "Email message",
                value=email_summary,
                height=300,
            )

    elif recipient:
        st.error(
            "Enter a valid email address."
        )

    st.subheader(
        "8. Calendar reminders"
    )

    st.download_button(
        "Download all calendar reminders",
        data=create_calendar_file(
            recurring
        ),
        file_name=(
            "subscription_reminders.ics"
        ),
        mime="text/calendar; charset=utf-8",
    )

    st.subheader(
        "9. Cancellation and dispute drafts"
    )

    if not enable_ai:
        st.info(
            "Enable optional AI drafts "
            "in the sidebar."
        )
    else:
        selected_merchant = st.selectbox(
            "Select a recurring charge",
            recurring["merchant"].tolist(),
        )

        selected_action = st.radio(
            "Draft type",
            [
                "Cancellation request",
                "Billing dispute note",
            ],
            horizontal=True,
        )

        selected = recurring[
            recurring["merchant"]
            == selected_merchant
        ].iloc[0]

        if st.button("Generate draft"):
            try:
                with st.spinner(
                    "Generating draft..."
                ):
                    draft = generate_ai_draft(
                        selected["merchant"],
                        selected["latest_amount"],
                        selected["previous_amount"],
                        selected["cadence"],
                        selected_action,
                    )

                st.text_area(
                    "Review before sending",
                    value=draft,
                    height=300,
                )

            except Exception as error:
                st.error(str(error))

    st.caption(
        "Verify all information against the "
        "original statements. This app does "
        "not provide financial or legal advice."
    )


if __name__ == "__main__":
    run_app()
